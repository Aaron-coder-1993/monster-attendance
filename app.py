import streamlit as st
import openpyxl
import json
import os
import re
from datetime import datetime, date
import calendar
import io
from copy import copy

# ================= 系統設定 =================
DATA_FILE = "MONSTER 隊員場地費.xlsx"  
HISTORY_FILE = "attendance_history.json" 

def get_role_from_fill(fill):
    """根據 Excel 儲存格顏色判斷身分"""
    if not fill or not fill.start_color:
        return "未分類"
    c = fill.start_color
    if c.type == 'theme':
        if c.theme == 4:
            return "🟦 底層"
        elif c.theme == 5:
            return "🟥 上層"
    return "未分類"

def load_history():
    if os.path.exists(HISTORY_FILE):
        with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def save_history(history):
    with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
        json.dump(history, f, ensure_ascii=False, indent=4)

def main():
    st.set_page_config(page_title="MONSTER 點名系統", page_icon="🏐", layout="wide")
    
    # ================= 手機版介面優化 =================
    st.markdown("""
        <style>
        .stMultiSelect span { font-size: 18px !important; }
        .stButton>button {
            width: 100%; height: 50px; font-size: 18px;
            font-weight: bold; border-radius: 8px;
        }
        </style>
    """, unsafe_allow_html=True)

    # ================= 1. 登入系統 =================
    if "logged_in" not in st.session_state:
        st.session_state["logged_in"] = False

    if not st.session_state["logged_in"]:
        st.title("🔒 MONSTER 點名系統 - 登入")
        st.markdown("---")
        col1, col2, col3 = st.columns([1, 2, 1])
        with col2:
            st.info("請輸入幹部帳號與密碼以進入系統")
            user_id = st.selectbox("👤 選擇點名帳號", ["A", "B", "C", "D", "E"])
            password = st.text_input("🔑 密碼", type="password")
            
            if st.button("✅ 登入系統"):
                if password == "MONSTER":
                    st.session_state["logged_in"] = True
                    st.session_state["current_user"] = user_id
                    st.rerun()
                else:
                    st.error("❌ 密碼錯誤！請重新輸入。")
        return 

    # ================= 已登入狀態 =================
    st.title("🏐 MONSTER 線上點名系統")
    col_title, col_logout = st.columns([8, 1])
    with col_title:
        st.caption(f"目前登入者：**{st.session_state['current_user']}** ｜ 對接檔案：`{DATA_FILE}`")
    with col_logout:
        if st.button("登出"):
            st.session_state["logged_in"] = False
            st.rerun()
    
    if "success_msg" in st.session_state:
        st.success(st.session_state["success_msg"])
        del st.session_state["success_msg"]

    if not os.path.exists(DATA_FILE):
        st.error(f"❌ 找不到檔案 `{DATA_FILE}`，請確認檔案已放入資料夾！")
        return
        
    try:
        wb = openpyxl.load_workbook(DATA_FILE)
    except Exception as e:
        st.error(f"❌ 檔案讀取失敗！請先關閉您的 Excel 再操作。詳細錯誤: {e}")
        return

    tab_attendance, tab_manage = st.tabs(["📝 點名作業", "⚙️ 名單管理"])

    # ================== 分頁 1：點名作業 ==================
    with tab_attendance:
        date_selected = st.date_input("📅 選擇點名日期", date.today())
        target_sheet_name = date_selected.strftime("%Y%m")
        
        sheet_exists = target_sheet_name in wb.sheetnames
        sheet_for_read = wb[target_sheet_name] if sheet_exists else wb[wb.sheetnames[-1]]

        members = []
        attendees_today = []
        claimed_tapes = [] # 紀錄本月已經領過白貼的人
        
        date_col = None
        tape_col = None
        current_taker = "無"
        
        guests_today = 0
        guest_row_idx = None
        purchased_tapes_today = 0
        purchase_row_idx = None
        
        if sheet_exists:
            # 找日期欄位
            for c in range(3, 34):
                c_val = sheet_for_read.cell(row=2, column=c).value
                if isinstance(c_val, datetime) and c_val.date() == date_selected:
                    date_col = c
                    if sheet_for_read.cell(row=1, column=c).value:
                        current_taker = str(sheet_for_read.cell(row=1, column=c).value)
                    break
            
            # 找「白貼」欄位 (通常在第34~40欄左右)
            for c in range(30, 45):
                if str(sheet_for_read.cell(row=3, column=c).value).strip() == "白貼":
                    tape_col = c
                    break

        # 解析名單與狀況
        for r in range(4, sheet_for_read.max_row + 1):
            val = sheet_for_read.cell(row=r, column=2).value
            if not val: continue
            name = str(val).strip()
            if not name: continue
            
            # 處理外賓與購買白貼專屬列
            if name == "外賓人數":
                guest_row_idx = r
                if date_col and sheet_for_read.cell(row=r, column=date_col).value:
                    try: guests_today = int(sheet_for_read.cell(row=r, column=date_col).value)
                    except: pass
                continue
            elif name == "購買白貼":
                purchase_row_idx = r
                if date_col and sheet_for_read.cell(row=r, column=date_col).value:
                    try: purchased_tapes_today = int(sheet_for_read.cell(row=r, column=date_col).value)
                    except: pass
                continue

            fill = sheet_for_read.cell(row=r, column=2).fill
            role = get_role_from_fill(fill)
            
            if role != "未分類": 
                members.append({"name": name, "role": role, "row": r})
                
                # 檢查當日出席
                if date_col and sheet_for_read.cell(row=r, column=date_col).value == 1:
                    attendees_today.append({"name": name, "role": role, "row": r})
                
                # 檢查本月是否已領白貼
                if tape_col and sheet_for_read.cell(row=r, column=tape_col).value:
                    claimed_tapes.append(name)

        # 排序
        history = load_history()
        bases = [m["name"] for m in members if m["role"] == "🟦 底層"]
        flyers = [m["name"] for m in members if m["role"] == "🟥 上層"]
        bases.sort(key=lambda x: history.get(x, 0), reverse=True)
        flyers.sort(key=lambda x: history.get(x, 0), reverse=True)
        
        # 準備白貼領取名單 (過濾掉已經領取過的人)
        available_for_tape = [m['name'] for m in members if m['name'] not in claimed_tapes]
        available_for_tape.sort(key=lambda x: history.get(x, 0), reverse=True)

        st.divider()
        col_left, col_right = st.columns([1.2, 1])

        with col_left:
            st.header("📝 新增點名與白貼登記")
            with st.form("attendance_form"):
                selected_bases = st.multiselect("🟦 選擇底層人員", options=bases)
                selected_flyers = st.multiselect("🟥 選擇上層人員", options=flyers)
                
                st.markdown("---")
                col_g, col_t = st.columns(2)
                with col_g:
                    guests_input_count = st.number_input("👤 當日外賓總人數", min_value=0, value=guests_today)
                with col_t:
                    tapes_purchased_input = st.number_input("📦 當日購買白貼數量", min_value=0, value=purchased_tapes_today)
                
                st.markdown("---")
                selected_tapes = st.multiselect("🩹 登記領取白貼 (已隱藏本月領過的人)", options=available_for_tape)
                
                if st.form_submit_button("✅ 送出並寫入 Excel"):
                    all_attendees = selected_bases + selected_flyers
                    try:
                        # 智慧換月功能
                        if not sheet_exists:
                            valid_sheets = [s for s in wb.sheetnames if re.match(r'^\d{6}$', s)]
                            valid_sheets.sort()
                            source_sheet_name = valid_sheets[-1] if valid_sheets else wb.sheetnames[-1]
                            source_sheet = wb[source_sheet_name]
                            
                            new_sheet = wb.copy_worksheet(source_sheet)
                            new_sheet.title = target_sheet_name
                            days_in_month = calendar.monthrange(date_selected.year, date_selected.month)[1]
                            
                            # 找出新表單的白貼欄位
                            new_tape_col = None
                            for c in range(30, 45):
                                if str(new_sheet.cell(row=3, column=c).value).strip() == "白貼":
                                    new_tape_col = c
                                    break
                            
                            for c in range(3, 34): 
                                day = c - 2
                                if day <= days_in_month:
                                    new_sheet.cell(row=2, column=c, value=datetime(date_selected.year, date_selected.month, day))
                                else:
                                    new_sheet.cell(row=2, column=c, value=None)
                                
                                new_sheet.cell(row=1, column=c, value=None) 
                                for r in range(4, new_sheet.max_row + 1):
                                    new_sheet.cell(row=r, column=c, value=None)
                            
                            # ✨ 換月時一併清空所有人的白貼紀錄
                            if new_tape_col:
                                for r in range(4, new_sheet.max_row + 1):
                                    new_sheet.cell(row=r, column=new_tape_col).value = None
                                    
                            target_sheet = new_sheet
                        else:
                            target_sheet = wb[target_sheet_name]

                        # 寫入目標工作表
                        target_col = None
                        target_tape_col = None
                        for c in range(3, 45):
                            c_val = target_sheet.cell(row=2, column=c).value
                            if isinstance(c_val, datetime) and c_val.date() == date_selected:
                                target_col = c
                            if str(target_sheet.cell(row=3, column=c).value).strip() == "白貼":
                                target_tape_col = c
                                
                        if target_col:
                            # 1. 寫入點名人員與打卡
                            target_sheet.cell(row=1, column=target_col).value = st.session_state["current_user"]
                            name_to_row = {str(target_sheet.cell(row=r, column=2).value).strip(): r for r in range(4, target_sheet.max_row + 1) if target_sheet.cell(row=r, column=2).value}
                                
                            for name in all_attendees:
                                if name in name_to_row:
                                    target_sheet.cell(row=name_to_row[name], column=target_col).value = 1
                                    
                            # 2. 寫入白貼領取 (打勾 V)
                            if target_tape_col:
                                for name in selected_tapes:
                                    if name in name_to_row:
                                        target_sheet.cell(row=name_to_row[name], column=target_tape_col).value = "V"
                            
                            # 3. 處理外賓與購買白貼專屬行
                            if not guest_row_idx:
                                guest_row_idx = target_sheet.max_row + 1
                                target_sheet.cell(row=guest_row_idx, column=1).value = "x"
                                target_sheet.cell(row=guest_row_idx, column=2).value = "外賓人數"
                            if guests_input_count > 0:
                                target_sheet.cell(row=guest_row_idx, column=target_col).value = guests_input_count
                            else:
                                target_sheet.cell(row=guest_row_idx, column=target_col).value = None
                                
                            if not purchase_row_idx:
                                purchase_row_idx = target_sheet.max_row + 1
                                target_sheet.cell(row=purchase_row_idx, column=1).value = "x"
                                target_sheet.cell(row=purchase_row_idx, column=2).value = "購買白貼"
                            if tapes_purchased_input > 0:
                                target_sheet.cell(row=purchase_row_idx, column=target_col).value = tapes_purchased_input
                            else:
                                target_sheet.cell(row=purchase_row_idx, column=target_col).value = None
                                    
                            wb.save(DATA_FILE)
                            
                            for p in all_attendees:
                                history[p] = history.get(p, 0) + 1
                            save_history(history)
                            
                            msg = f"🎉 成功儲存點名紀錄！負責人：{st.session_state['current_user']} "
                            if selected_tapes: msg += f" | 已登記 {len(selected_tapes)} 人領取白貼。"
                            st.session_state["success_msg"] = msg
                            st.rerun()
                        else:
                            st.error(f"❌ 找不到 {date_selected} 對應的日期欄位！")
                    except Exception as e:
                        st.error(f"❌ 發生錯誤: {e}")

        with col_right:
            st.header("📊 當日與本月概況")
            st.info(f"📝 本日點名負責人：**{current_taker}**")
            
            bases_present = [p['name'] for p in attendees_today if p['role'] == "🟦 底層"]
            flyers_present = [p['name'] for p in attendees_today if p['role'] == "🟥 上層"]
            
            st.markdown(f"**總出席隊員：{len(attendees_today)} 人**")
            st.markdown(f"🟦 **底層 ({len(bases_present)})：** " + "、".join(bases_present))
            st.markdown(f"🟥 **上層 ({len(flyers_present)})：** " + "、".join(flyers_present))
            
            col_s1, col_s2 = st.columns(2)
            col_s1.markdown(f"👤 **外賓人數：** `{guests_today}` 人")
            col_s2.markdown(f"📦 **購買白貼：** `{purchased_tapes_today}` 捲")
            
            st.write("---")
            st.markdown(f"🩹 **本月已領白貼人員 ({len(claimed_tapes)})**")
            if claimed_tapes:
                st.markdown("、".join(claimed_tapes))
            else:
                st.caption("本月尚無人領取")
            
            st.write("---")
            with st.form("delete_form"):
                st.markdown("⚠️ **發現登記錯誤？在這裡取消**")
                to_delete_attend = st.multiselect(
                    "取消「點名」 (可一次選多個)", 
                    options=[p['name'] for p in attendees_today]
                )
                to_delete_tape = st.multiselect(
                    "取消本月「白貼領取」", 
                    options=claimed_tapes
                )
                
                if st.form_submit_button("❌ 取消所選項並更新"):
                    if not to_delete_attend and not to_delete_tape:
                        st.warning("⚠️ 請先選擇要取消的項目！")
                    else:
                        try:
                            # 刪除出席
                            for name in to_delete_attend:
                                r = next(p['row'] for p in attendees_today if p['name'] == name)
                                sheet_for_read.cell(row=r, column=date_col).value = None
                            # 刪除白貼
                            if tape_col:
                                for name in to_delete_tape:
                                    r = next(m['row'] for m in members if m['name'] == name)
                                    sheet_for_read.cell(row=r, column=tape_col).value = None
                                    
                            wb.save(DATA_FILE)
                            st.session_state["success_msg"] = f"🗑️ 成功取消所選紀錄！"
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ 取消時發生錯誤: {e}")

    # ================== 分頁 2：名單管理 ==================
    with tab_manage:
        st.header("⚙️ 隊員名單增刪管理")
        st.write("在這裡新增、修改或刪除隊員，系統會自動幫您調整 Excel 表格與顏色配置。")
        col_add, col_edit, col_del = st.columns(3)
        
        with col_add:
            with st.form("add_member_form"):
                st.subheader("➕ 新增人員")
                new_name = st.text_input("輸入新隊員姓名")
                new_role = st.selectbox("選擇身分", ["🟦 底層", "🟥 上層"])
                
                if st.form_submit_button("新增至 Excel"):
                    new_name = new_name.strip()
                    if not new_name:
                        st.warning("請輸入姓名！")
                    elif new_name in [m['name'] for m in members]:
                        st.warning("此人員已經在名單中囉！")
                    else:
                        try:
                            insert_idx = None
                            ref_row = None
                            members_sorted = sorted(members, key=lambda x: x['row'])
                            
                            for m in reversed(members_sorted):
                                if m['role'] == new_role:
                                    insert_idx = m['row'] + 1
                                    ref_row = m['row']
                                    break
                            if not insert_idx:
                                insert_idx = 4
                                ref_row = 4
                                
                            sheet_for_read.insert_rows(insert_idx)
                            
                            for col in range(1, sheet_for_read.max_column + 1):
                                s_cell = sheet_for_read.cell(row=ref_row, column=col)
                                t_cell = sheet_for_read.cell(row=insert_idx, column=col)
                                
                                if s_cell.has_style:
                                    t_cell.font = copy(s_cell.font)
                                    t_cell.border = copy(s_cell.border)
                                    t_cell.fill = copy(s_cell.fill)
                                    t_cell.number_format = copy(s_cell.number_format)
                                    t_cell.alignment = copy(s_cell.alignment)
                                
                                if s_cell.data_type == 'f' and s_cell.value:
                                    old_formula = str(s_cell.value)
                                    pattern = r'([A-Z]+)' + str(ref_row)
                                    new_formula = re.sub(pattern, r'\g<1>' + str(insert_idx), old_formula)
                                    t_cell.value = new_formula

                            sheet_for_read.cell(row=insert_idx, column=1).value = "x" 
                            sheet_for_read.cell(row=insert_idx, column=2).value = new_name
                            
                            wb.save(DATA_FILE)
                            st.session_state["success_msg"] = f"✅ 成功將 {new_name} 加入 {new_role} 名單！"
                            st.rerun()
                        except Exception as e:
                            st.error(f"新增失敗：{e}")

        with col_edit:
            with st.form("edit_member_form"):
                st.subheader("✏️ 修改姓名")
                old_name = st.selectbox("選擇要修改的隊員", sorted([m['name'] for m in members]))
                edit_new_name = st.text_input("輸入正確/新的姓名")
                
                if st.form_submit_button("儲存修改"):
                    edit_new_name = edit_new_name.strip()
                    if not edit_new_name:
                        st.warning("請輸入新的姓名！")
                    elif edit_new_name == old_name:
                        st.warning("姓名沒有任何更動喔！")
                    elif edit_new_name in [m['name'] for m in members]:
                        st.warning("這個新名字已經存在名單中了！")
                    else:
                        try:
                            for m in members:
                                if m['name'] == old_name:
                                    sheet_for_read.cell(row=m['row'], column=2).value = edit_new_name
                                    break
                            
                            if old_name in history:
                                history[edit_new_name] = history.pop(old_name)
                            else:
                                history[edit_new_name] = 0
                            
                            wb.save(DATA_FILE)
                            save_history(history)
                            
                            st.session_state["success_msg"] = f"✏️ 成功將「{old_name}」修改為「{edit_new_name}」！"
                            st.rerun()
                        except Exception as e:
                            st.error(f"修改失敗：{e}")

        with col_del:
            with st.form("delete_member_form"):
                st.subheader("🗑️ 刪除人員")
                st.error("注意：刪除後會完全移除！")
                del_name = st.selectbox("選擇要永久刪除的隊員", sorted([m['name'] for m in members]))
                
                if st.form_submit_button("永久刪除該隊員"):
                    try:
                        for m in members:
                            if m['name'] == del_name:
                                sheet_for_read.delete_rows(m['row'], 1)
                                break
                        if del_name in history:
                            history.pop(del_name)
                            save_history(history)
                            
                        wb.save(DATA_FILE)
                        st.session_state["success_msg"] = f"🗑️ 已將 {del_name} 徹底移除！"
                        st.rerun()
                    except Exception as e:
                        st.error(f"刪除失敗：{e}")

    # ================== 備份下載 ==================
    st.divider()
    with open(DATA_FILE, "rb") as f:
        st.download_button(
            label="📥 下載最新的 Excel 檔案 (備份/查閱用)",
            data=f,
            file_name=f"MONSTER_{date_selected.strftime('%Y%m')}_點名表.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

if __name__ == "__main__":
    main()
